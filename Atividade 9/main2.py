from docx import Document

from docx.shared import Pt, Inches, RGBColor

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.enum.section import WD_SECTION

from docx.oxml import OxmlElement

from docx.oxml.ns import qn

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

from docx.enum.style import WD_STYLE_TYPE

from docx.enum.text import WD_BREAK

out = "/mnt/data/Caixas_de_Revisao_Apostila_Polimeros_QUI4715.docx"

doc = Document()

sec = doc.sections[0]

sec.top_margin = Inches(0.7)

sec.bottom_margin = Inches(0.7)

sec.left_margin = Inches(0.8)

sec.right_margin = Inches(0.8)

styles = doc.styles

styles["Normal"].font.name = "Arial"

styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph()

title.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = title.add_run("Proposta de Revisão Constante\nApostila de Polímeros – QUI 4715")

r.bold = True

r.font.size = Pt(18)

p = doc.add_paragraph()

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = p.add_run("Caixas de revisão para conectar novos conteúdos aos conhecimentos prévios")

r.italic = True

r.font.size = Pt(11)

doc.add_paragraph(

    "Objetivo geral: inserir momentos curtos de recuperação e conexão ao longo da apostila, "

    "evitando que o estudante tenha a sensação de estar aprendendo conteúdos novos de forma isolada. "

    "As atividades retomam principalmente ligações químicas, forças intermoleculares e funções orgânicas."

)

def shade_cell(cell, fill):

    tcPr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement('w:shd')

    shd.set(qn('w:fill'), fill)

    tcPr.append(shd)

def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):

    tc = cell._tc

    tcPr = tc.get_or_add_tcPr()

    tcMar = tcPr.first_child_found_in("w:tcMar")

    if tcMar is None:

        tcMar = OxmlElement('w:tcMar')

        tcPr.append(tcMar)

    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):

        node = tcMar.find(qn(f"w:{m}"))

        if node is None:

            node = OxmlElement(f"w:{m}")

            tcMar.append(node)

        node.set(qn("w:w"), str(v))

        node.set(qn("w:type"), "dxa")

def add_box(page, label, title, body_lines, objective, note=None):

    h = doc.add_paragraph()

    h.paragraph_format.space_before = Pt(10)

    h.paragraph_format.space_after = Pt(4)

    rr = h.add_run(f"Página {page} – {title}")

    rr.bold = True

    rr.font.size = Pt(13)

    

    table = doc.add_table(rows=1, cols=1)

    table.autofit = True

    cell = table.cell(0,0)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    shade_cell(cell, "EAF2F8")

    set_cell_margins(cell)

    

    p = cell.paragraphs[0]

    p.paragraph_format.space_after = Pt(5)

    r = p.add_run(label)

    r.bold = True

    r.font.size = Pt(11.5)

    

    for line in body_lines:

        p = cell.add_paragraph()

        p.paragraph_format.space_after = Pt(3)

        if isinstance(line, tuple):

            text, bold = line

            rr = p.add_run(text)

            rr.bold = bold

        else:

            p.add_run(line)

    

    p = doc.add_paragraph()

    p.paragraph_format.space_before = Pt(4)

    p.paragraph_format.space_after = Pt(2)

    r = p.add_run("Objetivo didático: ")

    r.bold = True

    p.add_run(objective)

    if note:

        p = doc.add_paragraph()

        p.paragraph_format.space_after = Pt(4)

        r = p.add_run("Observação: ")

        r.bold = True

        p.add_run(note)

boxes = [

("5","Relembrando...","Definição de polímero",

 ["Na revisão, vimos três tipos de ligações químicas: iônica, covalente e metálica.",

  "Os meros que formam uma cadeia polimérica estão unidos por ligações covalentes.",

  ("Você lembra? O que acontece com os elétrons em uma ligação covalente?", True),

  "☐ São transferidos entre os átomos.",

  "☐ São compartilhados entre os átomos.",

  "☐ Movimentam-se livremente pelo material."],

 "Retomar o conceito de ligação covalente antes de utilizá-lo na definição de polímero."),

("6","Observe...","PET e funções orgânicas",

 ["Observe novamente a estrutura do PET.",

  "Na revisão, você conheceu diferentes funções orgânicas.",

  ("Você consegue identificar na estrutura do PET uma função orgânica que contém oxigênio?", True),

  "Dica: procure o grupo –COO–.",

  "Função orgânica: ______________________________"],

 "Relacionar a estrutura química de um polímero a uma função orgânica já estudada.",

 "Resposta esperada: éster."),

("6","Conectando conceitos","Polimerização por adição",

 ["Lembra dos hidrocarbonetos?",

  "Na revisão, vimos que os alcenos apresentam uma ligação dupla entre carbonos.",

  "Observe agora o monômero do poliestireno.",

  ("1. Você consegue encontrar a ligação C=C? Circule-a na estrutura.", True),

  ("2. O que acontece com essa ligação durante a formação do polímero?", True),

  "____________________________________________________________"],

 "Recuperar o conceito de alceno e utilizá-lo para interpretar a polimerização por adição."),

("7","Observe e identifique","Formação do náilon 6,6",

 ["Você já conhece essas funções!",

  "Observe os dois monômeros utilizados na produção do náilon 6,6.",

  ("Hexametilenodiamina: qual função orgânica você identifica?", True),

  "________________________________",

  ("Ácido adípico: qual função orgânica você identifica?", True),

  "________________________________",

  "Agora observe o polímero formado.",

  ("Qual nova função orgânica aparece na cadeia do náilon?", True),

  "________________________________"],

 "Fazer o estudante reconhecer funções orgânicas diretamente nas estruturas usadas em uma reação de polimerização.",

 "Respostas esperadas: amina; ácido carboxílico; amida."),

("8","Conectando conceitos","Resistência do PET",

 ["Da molécula à propriedade.",

  "Na revisão, vimos que as moléculas podem interagir por meio de diferentes forças intermoleculares.",

  "Agora pense no PET:",

  ("Se as cadeias poliméricas apresentam interações mais fortes entre si, você espera que elas sejam:", True),

  "☐ mais fáceis de separar",

  "☐ mais difíceis de separar",

  ("Como isso pode ajudar a explicar a resistência mecânica da garrafa PET?", True),

  "____________________________________________________________",

  "____________________________________________________________"],

 "Relacionar interações entre cadeias com uma propriedade macroscópica do material."),

("8","Relembrando...","Borrachas e forças intermoleculares",

 [("As forças de Van der Waals são ligações dentro da cadeia ou interações entre diferentes cadeias?", True),

  "____________________________________________________________"],

 "Reforçar a distinção entre ligações covalentes na cadeia e interações intermoleculares entre cadeias."),

("9","Antes de continuar...","Estrutura e propriedades",

 ["Até aqui, você já viu que a estrutura de um polímero influencia suas propriedades.",

  ("1. Cadeias que interagem fortemente entre si tenderiam a ser mais ou menos fáceis de separar?", True),

  "____________________________________________________________",

  ("2. Cadeias que conseguem se movimentar com facilidade poderiam produzir um material mais rígido ou mais flexível?", True),

  "____________________________________________________________",

  "Agora compare suas respostas com as propriedades apresentadas a seguir."],

 "Estimular uma previsão antes da apresentação sistemática das propriedades dos polímeros."),

("11","Observe...","PLA e funções orgânicas",

 ["O PLA apareceu novamente!",

  "O PLA é utilizado em impressoras 3D e pode ser produzido a partir de matérias-primas renováveis.",

  "Observe a unidade de repetição do PLA.",

  ("Qual função orgânica estudada na revisão está presente na cadeia?", True),

  "☐ Álcool",

  "☐ Éster",

  "☐ Amina",

  "☐ Éter"],

 "Retomar funções orgânicas em um novo polímero e favorecer a transferência do conhecimento.",

 "Esta caixa requer que a estrutura/unidade de repetição do PLA seja inserida junto à atividade. Se a estrutura não for adicionada à apostila, recomenda-se retirar esta caixa."),

("12","Conectando conceitos","Recuperação acumulativa: PET e náilon",

 ["O que você já consegue reconhecer?",

  "Ao longo da apostila, você encontrou diferentes polímeros.",

  "PET → função ______________________________",

  "Náilon 6,6 → função ________________________",

  ("Esses grupos químicos fazem parte da cadeia do polímero ou são forças entre diferentes cadeias?", True),

  "____________________________________________________________"],

 "Verificar se o estudante consegue recuperar conceitos apresentados várias páginas antes e distinguir estrutura química de forças intermoleculares.",

 "Respostas esperadas: PET → éster; náilon 6,6 → amida."),

("12","Conectando conceitos","Termoplásticos e mobilidade das cadeias",

 ["Um termoplástico pode amolecer quando aquecido e ser moldado novamente.",

  ("Pensando no que você aprendeu sobre cadeias poliméricas, o que deve acontecer quando o material é aquecido?", True),

  "☐ As cadeias conseguem ganhar mobilidade.",

  "☐ Todas as ligações covalentes da cadeia são quebradas.",

  ("Qual alternativa parece mais coerente? Por quê?", True),

  "____________________________________________________________",

  "____________________________________________________________"],

 "Estimular raciocínio sobre o comportamento das cadeias ao relacionar estrutura e comportamento térmico.",

 "A mobilidade das cadeias não é desenvolvida detalhadamente antes na apostila. Portanto, esta questão deve ser tratada como atividade de raciocínio/discussão, e não como simples recuperação de conteúdo."),

("13","Conectando conceitos","Síntese final: do monômero ao produto",

 ["Escolha um dos materiais apresentados na apostila:",

  "PET | PVC | Poliestireno | Náilon 6,6",

  ("1. Polímero escolhido:", True),

  "________________________________",

  ("2. Onde ele aparece no dia a dia?", True),

  "____________________________________________________________",

  ("3. Que característica de sua estrutura você aprendeu?", True),

  "____________________________________________________________",

  ("4. Qual propriedade torna esse polímero útil nessa aplicação?", True),

  "____________________________________________________________",

  ("5. Que problema ambiental pode estar relacionado ao seu uso ou descarte?", True),

  "____________________________________________________________"],

 "Integrar estrutura, propriedade, aplicação e impacto ambiental em uma única atividade de fechamento.")

]

for b in boxes:

    add_box(*b)

doc.add_heading("Sugestão de padronização visual", level=1)

doc.add_paragraph(

    "Para manter a revisão constante sem poluir visualmente a apostila, recomenda-se utilizar apenas três categorias recorrentes:"

)

for name, desc in [

    ("Relembrando...", "recuperação rápida de um conceito já apresentado."),

    ("Observe...", "identificação de conceitos em estruturas, figuras ou exemplos."),

    ("Conectando conceitos", "aplicação de conhecimentos anteriores para compreender um conteúdo novo.")

]:

    p=doc.add_paragraph(style=None)

    p.style = doc.styles["Normal"]

    p.paragraph_format.left_indent = Inches(0.2)

    r=p.add_run(name + " ")

    r.bold=True

    p.add_run("– " + desc)

doc.add_paragraph(

    "A dificuldade das atividades pode aumentar progressivamente: recordar → reconhecer → identificar → explicar → prever → integrar. "

    "Assim, a revisão não funciona como repetição mecânica, mas como recuperação e reutilização do conhecimento."

)

doc.save(out)

print(out)